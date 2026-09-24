"""知识库文档加载。

目录名即文档类型，这个约定让切片策略可以自动适配：
    01-profile  -> profile   基本信息、个人档案
    02-projects -> projects  项目文档
    03-qa       -> qa        模拟面试问答对
    04-notes    -> notes     技术笔记
    05-jd       -> jd        目标岗位 JD
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from loguru import logger

from backend.config.settings import settings

# 目录前缀 -> 文档类型
DIR_TYPE_MAP = {
    "01-profile": "profile",
    "02-projects": "projects",
    "03-qa": "qa",
    "04-notes": "notes",
    "05-jd": "jd",
}

SUPPORTED_SUFFIX = {".md", ".markdown", ".txt", ".docx", ".pdf"}


@dataclass
class RawDocument:
    text: str
    source: str
    doc_type: str
    title: str
    metadata: dict = field(default_factory=dict)


def _read_markdown(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


_READERS = {
    ".md": _read_markdown,
    ".markdown": _read_markdown,
    ".txt": _read_markdown,
    ".docx": _read_docx,
    ".pdf": _read_pdf,
}


def _infer_doc_type(path: Path, root: Path) -> str:
    rel = path.relative_to(root)
    if len(rel.parts) > 1:
        return DIR_TYPE_MAP.get(rel.parts[0], "notes")
    return "notes"


def _infer_title(text: str, path: Path) -> str:
    for line in text.splitlines():
        line = line.strip()
        if line.startswith("# "):
            return line.lstrip("# ").strip()
    return path.stem


def load_documents(root: Path | None = None) -> list[RawDocument]:
    """递归加载知识库原始目录下的所有支持格式文档。"""
    root = root or settings.knowledge_raw_dir
    if not root.exists():
        logger.warning(f"知识库目录不存在: {root}")
        return []

    docs: list[RawDocument] = []
    files = sorted(
        p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIX
    )

    # 发布清单过滤：部署平台会复用持久卷，本地删除的文件在沙箱里仍然存在。
    # 若 knowledge/manifest.txt 存在，则只加载清单内的文件，避免上一版
    # 已删除 / 已脱敏的文档继续被检索到（公开部署下的信息泄露风险）。
    manifest = root.parent / "manifest.txt"
    if manifest.is_file():
        allowed = {
            line.strip().replace("\\", "/")
            for line in manifest.read_text(encoding="utf-8").splitlines()
            if line.strip() and not line.strip().startswith("#")
        }
        before = len(files)
        files = [
            p for p in files if str(p.relative_to(root)).replace("\\", "/") in allowed
        ]
        if before - len(files):
            logger.info(f"按发布清单忽略 {before - len(files)} 个残留文件")

    for path in files:
        try:
            text = _READERS[path.suffix.lower()](path)
        except Exception as e:
            logger.warning(f"读取失败，已跳过 {path.name}: {e}")
            continue

        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        if not text:
            continue

        docs.append(
            RawDocument(
                text=text,
                source=str(path.relative_to(root)),
                doc_type=_infer_doc_type(path, root),
                title=_infer_title(text, path),
            )
        )

    logger.info(f"共加载 {len(docs)} 篇文档，来自 {root}")
    return docs


def write_manifest(raw_dir: Path | None = None) -> Path | None:
    """把 raw 目录下的文档相对路径写入 manifest.txt（发布清单）。

    供 sync_deploy（发布前）与 bootstrap（写入新资料后）调用：
    清单必须与「当前 raw 目录」严格一致，否则新写入的文件会被
    load_documents() 的清单过滤挡掉，表现为「上传成功但检索不到」。
    """
    root = raw_dir or settings.knowledge_raw_dir
    if not root.exists():
        return None
    entries = sorted(
        str(p.relative_to(root)).replace("\\", "/")
        for p in root.rglob("*")
        if p.is_file() and p.suffix.lower() != ".pyc"
    )
    manifest = root.parent / "manifest.txt"
    manifest.write_text("\n".join(entries) + "\n", encoding="utf-8")
    logger.info(f"已更新发布清单 {manifest}（{len(entries)} 个文档）")
    return manifest
